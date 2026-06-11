import os

import chromadb
from docx import Document
from pypdf import PdfReader
from sentence_transformers import SentenceTransformer

from .models import KnowledgeDocument


CHROMA_PATH = "chroma_db"
COLLECTION_NAME = "localmind_knowledge"

embedding_model = None

client = chromadb.PersistentClient(path=CHROMA_PATH)
collection = client.get_or_create_collection(name=COLLECTION_NAME)


def get_embedding_model():
    global embedding_model

    if embedding_model is None:
        embedding_model = SentenceTransformer("all-MiniLM-L6-v2")

    return embedding_model


def chunk_text(text, chunk_size=500):
    words = text.split()
    chunks = []

    for i in range(0, len(words), chunk_size):
        chunk = " ".join(words[i:i + chunk_size])

        if chunk.strip():
            chunks.append(chunk)

    return chunks


def extract_text_from_pdf_with_pypdf(file_path):
    reader = PdfReader(file_path)
    text_parts = []

    for page in reader.pages:
        page_text = page.extract_text() or ""
        if page_text.strip():
            text_parts.append(page_text.strip())

    return "\n".join(text_parts)


def extract_text_from_pdf_with_ocr(file_path):
    """
    OCR fallback for scanned/image PDFs.
    Requires:
    - pytesseract
    - pdf2image
    - pillow
    - Tesseract OCR installed on Windows
    - Poppler installed on Windows
    """
    try:
        import pytesseract
        from pdf2image import convert_from_path
    except ImportError:
        return ""

    poppler_path = os.getenv("POPPLER_PATH", "").strip()
    tesseract_cmd = os.getenv("TESSERACT_CMD", "").strip()

    if tesseract_cmd:
        pytesseract.pytesseract.tesseract_cmd = tesseract_cmd

    try:
        if poppler_path:
            images = convert_from_path(file_path, dpi=220, poppler_path=poppler_path)
        else:
            images = convert_from_path(file_path, dpi=220)

        text_parts = []

        for image in images:
            page_text = pytesseract.image_to_string(image, lang="eng") or ""
            if page_text.strip():
                text_parts.append(page_text.strip())

        return "\n".join(text_parts)

    except Exception as error:
        print(f"OCR PDF extraction failed: {error}")
        return ""


def extract_text_from_pdf(file_path):
    text = extract_text_from_pdf_with_pypdf(file_path)

    if len(text.strip()) >= 30:
        return text

    print("PDF text extraction returned very little text. Trying OCR fallback...")
    ocr_text = extract_text_from_pdf_with_ocr(file_path)

    if ocr_text.strip():
        return ocr_text

    return text


def extract_text_from_txt(file_path):
    with open(file_path, "r", encoding="utf-8", errors="ignore") as file:
        return file.read()


def extract_text_from_docx(file_path):
    document = Document(file_path)
    text_parts = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text_parts.append(paragraph.text.strip())

    for table in document.tables:
        for row in table.rows:
            row_text = []

            for cell in row.cells:
                if cell.text.strip():
                    row_text.append(cell.text.strip())

            if row_text:
                text_parts.append(" | ".join(row_text))

    return "\n".join(text_parts)


def extract_text_from_document(file_path):
    lower_path = file_path.lower()

    if lower_path.endswith(".pdf"):
        return extract_text_from_pdf(file_path)

    if lower_path.endswith(".txt"):
        return extract_text_from_txt(file_path)

    if lower_path.endswith(".docx"):
        return extract_text_from_docx(file_path)

    return ""


def clear_user_collection(user):
    existing = collection.get(where={"user_id": user.id})

    if existing and existing.get("ids"):
        collection.delete(ids=existing["ids"])


def build_knowledge_base(user):
    clear_user_collection(user)

    documents = KnowledgeDocument.objects.filter(user=user)
    chunk_count = 0
    model = get_embedding_model()

    for document in documents:
        file_path = document.file.path

        if not os.path.exists(file_path):
            continue

        text = extract_text_from_document(file_path)

        if not text.strip():
            continue

        chunks = chunk_text(text)

        for index, chunk in enumerate(chunks):
            chunk_id = f"user_{user.id}_doc_{document.id}_chunk_{index}"
            embedding = model.encode(chunk).tolist()

            collection.add(
                ids=[chunk_id],
                documents=[chunk],
                embeddings=[embedding],
                metadatas=[
                    {
                        "user_id": user.id,
                        "document_id": document.id,
                        "source": document.title,
                        "file_name": os.path.basename(file_path),
                    }
                ],
            )

            chunk_count += 1

    return chunk_count


def search_knowledge(query, user, top_k=3):
    model = get_embedding_model()
    query_embedding = model.encode(query).tolist()

    results = collection.query(
        query_embeddings=[query_embedding],
        n_results=top_k,
        where={"user_id": user.id},
    )

    documents = results.get("documents", [[]])[0]
    metadatas = results.get("metadatas", [[]])[0]

    context_parts = []

    for index, document in enumerate(documents):
        if not document or not document.strip():
            continue

        metadata = metadatas[index] if index < len(metadatas) else {}
        source = metadata.get("source", "unknown")

        context_parts.append(f"Source: {source}\nContent: {document}")

    return "\n\n".join(context_parts)


def search_knowledge_with_sources(query, user, top_k=3):
    model = get_embedding_model()
    query_embedding = model.encode(query).tolist()

    results = collection.query(
        query_embeddings=[query_embedding],
        n_results=top_k,
        where={"user_id": user.id},
    )

    documents = results.get("documents", [[]])[0]
    metadatas = results.get("metadatas", [[]])[0]

    context_parts = []
    sources = []

    for index, document in enumerate(documents):
        if not document or not document.strip():
            continue

        metadata = metadatas[index] if index < len(metadatas) else {}

        source = metadata.get("source", "unknown")
        document_id = metadata.get("document_id")
        file_name = metadata.get("file_name", "unknown")

        context_parts.append(f"Source: {source}\nContent: {document}")

        sources.append(
            {
                "document_id": document_id,
                "source": source,
                "file_name": file_name,
                "chunk_preview": document[:300],
            }
        )

    context = "\n\n".join(context_parts)

    return {
        "has_context": bool(context.strip()),
        "context": context,
        "sources": sources,
    }


def get_user_collection_stats(user):
    existing = collection.get(where={"user_id": user.id})
    ids = existing.get("ids", []) if existing else []

    return {
        "total_chunks": len(ids),
        "is_indexed": len(ids) > 0,
    }


def get_document_chunks(user, document_id):
    results = collection.get(
        where={
            "$and": [
                {"user_id": user.id},
                {"document_id": document_id},
            ]
        }
    )

    chunks = []

    documents = results.get("documents", []) if results else []
    metadatas = results.get("metadatas", []) if results else []
    ids = results.get("ids", []) if results else []

    for index, document in enumerate(documents):
        metadata = metadatas[index] if index < len(metadatas) else {}

        chunks.append(
            {
                "chunk_id": ids[index] if index < len(ids) else "",
                "source": metadata.get("source"),
                "file_name": metadata.get("file_name"),
                "content": document,
                "preview": document[:300],
            }
        )

    return chunks


def rebuild_single_document(user, document):
    existing = collection.get(
        where={
            "$and": [
                {"user_id": user.id},
                {"document_id": document.id},
            ]
        }
    )

    if existing and existing.get("ids"):
        collection.delete(ids=existing["ids"])

    file_path = document.file.path

    if not os.path.exists(file_path):
        return 0

    text = extract_text_from_document(file_path)

    if not text.strip():
        return 0

    chunks = chunk_text(text)
    model = get_embedding_model()
    chunk_count = 0

    for index, chunk in enumerate(chunks):
        chunk_id = f"user_{user.id}_doc_{document.id}_chunk_{index}"
        embedding = model.encode(chunk).tolist()

        collection.add(
            ids=[chunk_id],
            documents=[chunk],
            embeddings=[embedding],
            metadatas=[
                {
                    "user_id": user.id,
                    "document_id": document.id,
                    "source": document.title,
                    "file_name": os.path.basename(file_path),
                }
            ],
        )

        chunk_count += 1

    return chunk_count


def delete_document_from_vector_db(user, document_id):
    existing = collection.get(
        where={
            "$and": [
                {"user_id": user.id},
                {"document_id": document_id},
            ]
        }
    )

    ids = existing.get("ids", []) if existing else []

    if ids:
        collection.delete(ids=ids)

    return len(ids)