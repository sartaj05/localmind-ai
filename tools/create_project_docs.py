from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"

NAVY = "17365D"
BLUE = "1F4E79"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "F3F8FC"
GRAY = "666666"
LIGHT_GRAY = "D9E1F2"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    p_pr.append(keep)


def set_font(run, name="Aptos", size=10.5, color="222222", bold=False, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def configure_document(doc, short_title):
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("222222")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for style_name, size, color in (("Title", 24, NAVY), ("Heading 1", 16, NAVY), ("Heading 2", 12.5, BLUE), ("Heading 3", 11, BLUE)):
        style = styles[style_name]
        style.font.name = "Aptos Display" if style_name == "Title" else "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(f"{short_title}  |  LocalMind AI")
    set_font(run, size=8, color=GRAY)


def add_title(doc, title, subtitle):
    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(title)
    set_font(run, name="Aptos Display", size=24, color=NAVY, bold=True)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(12)
    run = p2.add_run(subtitle)
    set_font(run, size=11, color=GRAY, italic=True)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    set_keep_with_next(p)
    return p


def add_para(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_font(r, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_font(r2)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_font(r, size=10.2)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        if widths:
            cell.width = Inches(widths[i])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(text)
        set_font(r, size=9.2, color="FFFFFF", bold=True)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cell = cells[i]
            if widths:
                cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            r = p.add_run(str(text))
            set_font(r, size=9.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_status_line(doc, label, text, color=BLUE):
    p = doc.add_paragraph()
    r = p.add_run(label + " ")
    set_font(r, bold=True, color=color)
    r2 = p.add_run(text)
    set_font(r2)


def save(doc, filename):
    path = DOCS / filename
    doc.save(path)
    return path


def feature_inventory():
    doc = Document()
    configure_document(doc, "Feature Inventory")
    add_title(doc, "LocalMind AI Feature Inventory", "Current product capability, launch readiness, and technical scope")
    add_para(doc, "This inventory describes the current repository as of the latest committed project state. It separates implemented capabilities from partially wired work so the document can be used in a portfolio review, launch discussion, or technical interview.")
    add_status_line(doc, "Launch recommendation:", "Suitable for a portfolio demo or controlled beta after a short hardening pass. It is not yet ready for an unrestricted public production launch.", color="9C0006")

    add_heading(doc, "Product Summary", 1)
    add_para(doc, "LocalMind AI is a privacy-focused full-stack AI workspace. It combines a React interface with a Django REST API, local Ollama model inference, retrieval-augmented generation for uploaded documents, chat lifecycle management, prompt reuse, analytics, and the initial data model for collaborative workspaces.")

    add_heading(doc, "Implemented Features", 1)
    rows = [
        ("Authentication", "Registration, profile, JWT login, refresh, authenticated API permissions", "Implemented"),
        ("AI chat", "Local Ollama model calls, model listing, health check, chat history, chat sessions, message persistence", "Implemented"),
        ("Streaming", "Plain responses plus streaming endpoints, including SSE-oriented preference support", "Implemented"),
        ("Chat organization", "Search, rename, pin, archive, trash, restore, permanent delete, duplicate, merge, tags, important messages, bulk actions", "Implemented"),
        ("RAG knowledge base", "PDF, TXT, and DOCX extraction; embeddings; Chroma persistence; source-aware answers; chunk preview; rebuild and vector deletion", "Implemented"),
        ("OCR fallback", "Scanned PDF fallback through Tesseract and Poppler configuration", "Implemented with external prerequisites"),
        ("Exports", "TXT and JSON export for sessions and knowledge history, including bulk export", "Implemented"),
        ("Prompt library", "Create, edit, delete, search, categorize, pin, and use reusable prompts", "Implemented"),
        ("Usage controls", "Per-user daily limits, character limits, throttling, usage logs, daily usage summaries", "Implemented"),
        ("Analytics", "Summary cards plus daily requests, model usage, and success-rate charts", "Implemented"),
        ("AI preferences", "Default model, RAG top-k, source visibility, auto titles, stream format, quota, and prompt limits", "Implemented"),
        ("Collaboration", "Workspace, member, invitation data models and admin registration", "Started; API and UI are not wired"),
        ("API documentation", "OpenAPI schema and Swagger UI endpoints", "Implemented"),
    ]
    add_table(doc, ["Area", "Current capability", "Status"], rows, widths=[1.35, 4.55, 1.1])

    add_heading(doc, "Current Frontend Experience", 1)
    add_bullets(doc, [
        "React and Vite single-page interface with login and registration flows.",
        "Chat workspace with session navigation, model selection, streaming answers, message actions, archive and trash views.",
        "Knowledge workspace for upload, search, document inspection, chunk previews, rebuild, and RAG questions with sources.",
        "Usage dashboard with quota cards and Recharts visualizations.",
        "Settings for AI preferences and prompt library for reusable instructions.",
    ])

    add_heading(doc, "Launch Readiness Assessment", 1)
    add_para(doc, "The core product loop is strong enough to demonstrate: authenticate, create a conversation, ask a local model, upload documents, ask grounded questions, review sources, organize history, and inspect usage. The launch gap is mainly production engineering and completeness rather than lack of product direction.")
    rows = [
        ("P0", "Database migrations", "Generate and commit migrations for accounts, ai_engine, and collaboration; run migrate from a clean database", "Required before launch"),
        ("P0", "Security configuration", "Production secret, DEBUG false, strict hosts and CORS, HTTPS, secure cookies, upload limits", "Required before public launch"),
        ("P0", "Testing", "Replace placeholder tests with API, permission, RAG, streaming, export, and regression coverage", "Required before public launch"),
        ("P0", "Deployment", "Add production settings, process manager, persistent media and Chroma storage, backups, and monitoring", "Required before public launch"),
        ("P1", "Collaboration delivery", "Wire workspace endpoints, invitations, permissions, frontend screens, and migrations", "Required for team positioning"),
        ("P1", "Reliability", "Handle Ollama timeouts, model-unavailable states, RAG indexing jobs, retries, and user-facing errors", "Strongly recommended"),
        ("P1", "Product polish", "Onboarding, empty states, keyboard accessibility, responsive QA, file validation, and confirmation flows", "Strongly recommended"),
        ("P2", "Scale", "PostgreSQL, Redis/Celery jobs, object storage, observability, and usage-cost reporting", "Future"),
    ]
    add_table(doc, ["Priority", "Area", "Action", "Decision"], rows, widths=[0.55, 1.4, 4.15, 0.9])

    add_heading(doc, "Technology Stack", 1)
    rows = [
        ("Frontend", "React 19, Vite, React Router, Axios, Recharts, Lucide React"),
        ("Backend", "Python, Django 5.2, Django REST Framework, SimpleJWT, drf-spectacular"),
        ("AI runtime", "Ollama local inference with configurable model and base URL"),
        ("RAG", "Sentence Transformers all-MiniLM-L6-v2, ChromaDB, pypdf, python-docx, optional Tesseract and Poppler"),
        ("Data", "SQLite for development; production database should be PostgreSQL"),
        ("Runtime", "Django ASGI/WSGI, Uvicorn available, local Vite development server"),
    ]
    add_table(doc, ["Layer", "Technologies"], rows, widths=[1.4, 5.6])

    add_heading(doc, "Honest Portfolio Positioning", 1)
    add_para(doc, "Present LocalMind AI as a working local-first AI assistant MVP with a production-minded roadmap. Demonstrate the completed chat and RAG flows live. Describe collaboration as an in-progress domain model, not as a completed team feature, until its API, migrations, permissions, and UI are shipped.")
    add_heading(doc, "Recommended Launch Decision", 1)
    add_para(doc, "Launch a private demo or invite-only beta after P0 hardening. Delay a public production launch until clean migrations, security settings, tests, deployment persistence, and Ollama failure handling are complete.")
    return save(doc, "LocalMind_AI_Feature_Inventory.docx")


def roadmap():
    doc = Document()
    configure_document(doc, "Product Roadmap")
    add_title(doc, "LocalMind AI Product Roadmap", "Upcoming features and an implementation order for launch and growth")
    add_para(doc, "This roadmap turns the current LocalMind AI codebase into a staged product plan. It is ordered by risk reduction first, then by user value, collaboration, and scale. Each phase has a clear exit condition so new features do not hide unfinished launch work.")

    add_heading(doc, "Phase Zero Production Hardening", 1)
    add_para(doc, "Goal: make the existing MVP safe and repeatable to install, test, and operate.")
    add_bullets(doc, [
        "Generate and commit Django migrations for every app, including the new collaboration app.",
        "Split development and production settings; remove the fallback secret and set DEBUG false in production.",
        "Add upload type, size, filename, and content validation; protect media files from unintended public exposure.",
        "Add automated tests for auth, ownership, session actions, RAG isolation, exports, prompts, preferences, and quota enforcement.",
        "Add Ollama timeout, unavailable-model, empty-response, and indexing error states with retry-safe behavior.",
        "Persist media and Chroma data with backups; document restore and cleanup procedures.",
    ])
    add_status_line(doc, "Exit condition:", "A new developer can install the project from the deployment guide, run migrations from an empty database, pass the test suite, and operate the core demo flow.")

    add_heading(doc, "Phase One Collaboration Release", 1)
    add_para(doc, "Goal: complete the collaboration concept already represented in the repository.")
    add_table(doc, ["Feature", "Implementation", "Acceptance signal"], [
        ("Workspace APIs", "List, create, detail, update, delete with membership scoping", "Member sees only permitted workspaces"),
        ("Membership", "Add/remove members, roles, permission checks, audit events", "Admin/editor/viewer behavior is test-covered"),
        ("Invitations", "Create, expire, accept, reject, revoke, resend", "Token cannot be reused or accepted after expiry"),
        ("Shared context", "Workspace-scoped sessions, prompts, documents, and dashboards", "Users can collaborate without cross-tenant leakage"),
        ("Frontend", "Workspace switcher, members page, invite dialog, role management", "A user can complete the collaboration flow without API tools"),
    ], widths=[1.45, 4.25, 1.3])

    add_heading(doc, "Phase Two Product Experience", 1)
    add_bullets(doc, [
        "Onboarding that checks Ollama availability, confirms a model, and creates a first prompt.",
        "Conversation folders, saved views, keyboard shortcuts, message copy feedback, and richer markdown/code rendering.",
        "Document ingestion status, per-document indexing progress, duplicate detection, and unsupported-file feedback.",
        "Prompt variables, prompt version history, favorites, and shareable workspace prompt collections.",
        "Accessibility and responsive pass: keyboard navigation, labels, focus states, contrast, mobile layouts, and screen-reader semantics.",
    ])

    add_heading(doc, "Phase Three Production Platform", 1)
    add_bullets(doc, [
        "PostgreSQL for application data and Redis plus Celery for document indexing and title generation jobs.",
        "Object storage for user files, signed downloads, retention policies, and virus scanning.",
        "Structured logs, request IDs, metrics, error tracking, health/readiness probes, and alerting.",
        "Rate limits by user and workspace, model concurrency limits, and administrator usage reports.",
        "CI pipeline for lint, tests, migrations check, frontend build, security scan, and deployment artifact creation.",
    ])

    add_heading(doc, "Phase Four Intelligence and Differentiation", 1)
    add_bullets(doc, [
        "Multiple model profiles with per-task defaults and model comparison.",
        "Hybrid retrieval with metadata filters, reranking, citations, and configurable chunk strategies.",
        "Document collections, incremental indexing, and conversation-aware retrieval.",
        "Evaluation set for groundedness, citation quality, latency, and refusal behavior.",
        "Optional cloud model adapters while retaining the local-first privacy mode.",
    ])

    add_heading(doc, "Priority Matrix", 1)
    add_table(doc, ["Priority", "Ship next", "Why"], [
        ("Now", "Migrations, production settings, tests, deployment guide, failure handling", "Removes launch blockers"),
        ("Next", "Workspace APIs and UI, onboarding, indexing status, accessibility", "Turns the MVP into a usable product"),
        ("Later", "PostgreSQL, background jobs, object storage, observability", "Improves reliability and scale"),
        ("Explore", "Hybrid retrieval, evaluations, model adapters", "Creates long-term differentiation"),
    ], widths=[1.0, 3.8, 2.2])

    add_heading(doc, "What Not to Add Yet", 1)
    add_para(doc, "Avoid adding many model providers, billing, a mobile app, or complex agent workflows before the P0 foundation is stable. The current product already has enough surface area for a credible launch story; reliability, security, and complete collaboration will create more value than another large feature branch.")
    return save(doc, "LocalMind_AI_Product_Roadmap.docx")


def interview_prep():
    doc = Document()
    configure_document(doc, "Interview Preparation")
    add_title(doc, "LocalMind AI Interview Preparation", "A practical discussion guide for a developer with two or more years of experience")
    add_para(doc, "Use this guide to explain the project accurately and confidently. The strongest interview narrative is a clear product problem, a simple architecture, a few meaningful engineering decisions, and an honest account of what is complete versus next.")

    add_heading(doc, "Thirty Second Project Introduction", 1)
    add_para(doc, "LocalMind AI is a privacy-focused AI workspace built with React and Django REST Framework. Users authenticate with JWT, chat with local Ollama models, upload PDF, TXT, or DOCX documents, and ask grounded questions through a ChromaDB-backed RAG pipeline. The application also provides streaming responses, chat organization, reusable prompts, exports, usage limits, and analytics. I designed it as a local-first system so sensitive documents can remain on the user's machine or controlled infrastructure.")

    add_heading(doc, "Architecture in Brief", 1)
    add_table(doc, ["Layer", "What to explain"], [
        ("React client", "Single-page UI calls the API through Axios, stores short-lived UI state, and renders chat, RAG, prompt, settings, and dashboard panels."),
        ("Django API", "REST endpoints enforce JWT authentication and user ownership, persist sessions/messages/documents/preferences, and expose OpenAPI documentation."),
        ("AI service", "A service layer calls Ollama for local generation, model discovery, health checks, title generation, and streaming."),
        ("RAG service", "Documents are extracted, chunked, embedded with Sentence Transformers, stored in Chroma, then retrieved with user scoping before prompt construction."),
        ("Persistence", "SQLite is used for development; media stores uploads and Chroma stores vectors. Production should move application data to PostgreSQL and use durable storage."),
    ], widths=[1.4, 5.6])

    add_heading(doc, "Technology Stack in Brief", 1)
    add_table(doc, ["Technology", "Role"], [
        ("React 19 and Vite", "Frontend component model and local development/build workflow"),
        ("Django 5.2 and DRF", "Backend, ORM, API views, serializers, permissions, and admin"),
        ("SimpleJWT", "Access and refresh token authentication"),
        ("Ollama", "Local LLM inference and model discovery"),
        ("Sentence Transformers", "Text embeddings for semantic retrieval"),
        ("ChromaDB", "Persistent vector collection and similarity search"),
        ("SQLite, Axios, Recharts", "Development database, HTTP client, and analytics charts"),
    ], widths=[2.0, 5.0])

    add_heading(doc, "Likely Technical Questions and Answer Points", 1)
    questions = [
        ("Why local AI?", "Privacy, offline control, predictable ownership of documents, and no mandatory provider cost. The trade-off is local hardware, model management, and performance variability."),
        ("How does RAG work here?", "Extract supported document text, split it into chunks, embed chunks, store metadata and vectors, retrieve top-k chunks scoped to the user, and place the context in a guarded prompt before Ollama generation."),
        ("How do you prevent cross-user retrieval?", "The vector metadata stores user_id and queries include the current user filter. Django document and session querysets also use the authenticated user. I would add explicit isolation tests before production."),
        ("Why streaming?", "It reduces perceived latency by showing tokens as they arrive. The backend supports streaming responses, while the client maintains the partial assistant message and handles completion or failure."),
        ("How would you scale indexing?", "Move embedding and extraction to background jobs with Celery and Redis, track document states, make jobs idempotent, and persist vectors in a managed or durable vector service."),
        ("What are the current production risks?", "Migrations are not currently committed, settings are development-oriented, tests are minimal, SQLite is not a production database, and collaboration is only partly wired. I would resolve those before public launch."),
        ("How do you secure the API?", "JWT authentication, ownership filters, permission checks, throttling, input validation, strict CORS/hosts, secure production settings, upload validation, and audit logging. The last four still need production hardening."),
        ("Why use serializers?", "They centralize validation and representation, prevent accidental field exposure, and keep API contracts explicit."),
        ("How would you test RAG?", "Unit-test extraction and chunking, integration-test indexing and retrieval, verify user isolation, test empty and scanned documents, and use a small evaluation set for grounded answers and citations."),
        ("What would you change first?", "Commit migrations, add a production settings split, add high-value API tests, make indexing asynchronous, and complete workspace permissions and UI."),
    ]
    for question, answer in questions:
        p = doc.add_paragraph()
        r = p.add_run(question)
        set_font(r, size=10.5, color=BLUE, bold=True)
        p.paragraph_format.space_after = Pt(2)
        add_para(doc, answer)

    add_heading(doc, "Behavioral Questions to Prepare", 1)
    add_bullets(doc, [
        "Tell me about a difficult bug: use a real example and explain diagnosis, evidence, fix, and regression prevention.",
        "Tell me about a trade-off: discuss local inference versus hosted APIs or SQLite versus PostgreSQL.",
        "How do you handle an incomplete feature: say what is implemented, what is missing, how you document it, and how you prioritize the next slice.",
        "How do you review code: check correctness, security, tests, API compatibility, and operational impact before style preferences.",
        "How do you learn a new tool: build a small vertical slice, read primary documentation, measure behavior, and record decisions.",
    ])

    add_heading(doc, "Demonstration Flow", 1)
    add_bullets(doc, [
        "Register or log in and explain JWT access and refresh tokens.",
        "Create a chat, select a locally available model, send a prompt, and show streaming output.",
        "Create a prompt template and reuse it in chat.",
        "Upload a document, build or rebuild its knowledge index, ask a grounded question, and show sources.",
        "Open the dashboard to explain quotas, request counts, model usage, and success rate.",
        "Close with the production roadmap and clearly label collaboration as the next feature if the API/UI is not yet connected.",
    ])

    add_heading(doc, "Interview Accuracy Checklist", 1)
    add_table(doc, ["Say", "Avoid saying"], [
        ("Working local-first AI assistant MVP", "Production-ready SaaS platform"),
        ("RAG pipeline with user-scoped vector metadata", "Perfectly secure multi-tenant RAG"),
        ("Collaboration domain model started", "Complete real-time collaboration"),
        ("SQLite for development and a PostgreSQL migration plan", "SQLite is the final production database"),
        ("Streaming endpoints and UI support", "Guaranteed token streaming under all failures"),
    ], widths=[3.5, 3.5])
    add_heading(doc, "Closing Statement", 1)
    add_para(doc, "The project demonstrates practical full-stack experience across authentication, API design, local AI integration, RAG, asynchronous-looking streaming flows, data modeling, UI state, analytics, and operational thinking. For a two-plus-year interview, emphasize how you made decisions, verified behavior, handled trade-offs, and identified the next engineering risks.")
    return save(doc, "LocalMind_AI_Interview_Preparation.docx")


def deployment_reference():
    doc = Document()
    configure_document(doc, "Deployment Reference")
    add_title(doc, "LocalMind AI Deployment Reference", "Local setup, production preparation, environment variables, and operational checks")
    add_para(doc, "This reference is written for the current repository. It covers a Windows development setup and a production-shaped deployment plan. The repository is currently configured for SQLite and local Ollama, so public deployment requires the production hardening items called out below.")
    add_heading(doc, "Prerequisites", 1)
    add_bullets(doc, [
        "Python 3.11 or newer with a working virtual environment.",
        "Node.js 18 or newer and npm for the Vite frontend.",
        "Ollama installed and running locally or on a reachable private host.",
        "An Ollama model pulled before first use, for example: ollama pull phi3.",
        "Optional OCR support: Tesseract OCR and Poppler; configure their executable paths for scanned PDFs.",
        "For public production: PostgreSQL, a reverse proxy with HTTPS, durable media storage, and durable Chroma/vector storage or a managed vector service.",
    ])
    add_heading(doc, "Environment Configuration", 1)
    add_table(doc, ["Variable", "Example", "Purpose"], [
        ("SECRET_KEY", "replace-with-random-secret", "Django signing and security key"),
        ("DEBUG", "True for local; False for production", "Debug mode switch"),
        ("ALLOWED_HOSTS", "127.0.0.1,localhost", "Allowed host names"),
        ("OLLAMA_BASE_URL", "http://localhost:11434", "Ollama API address"),
        ("DEFAULT_AI_MODEL", "phi3", "Default local model"),
        ("POPPLER_PATH", "optional Windows path", "Poppler path for OCR PDF conversion"),
        ("TESSERACT_CMD", "optional tesseract.exe path", "Tesseract executable for OCR"),
    ], widths=[1.75, 2.25, 3.0])

    add_heading(doc, "Windows Local Setup", 1)
    add_table(doc, ["Step", "Command or action"], [
        ("1", "Open PowerShell in the repository root."),
        ("2", "Create an environment: python -m venv venv"),
        ("3", "Activate it: .\\venv\\Scripts\\Activate.ps1"),
        ("4", "Install backend packages: python -m pip install -r requirements.txt"),
        ("5", "Create .env from .env.example and set a development SECRET_KEY."),
        ("6", "Generate local migrations: python manage.py makemigrations"),
        ("7", "Apply migrations: python manage.py migrate"),
        ("8", "Create an admin user: python manage.py createsuperuser"),
        ("9", "Start Django: python manage.py runserver"),
        ("10", "In frontend, run npm install and then npm run dev."),
        ("11", "Open the Vite URL, normally http://localhost:5173."),
    ], widths=[0.55, 6.45])
    add_para(doc, "Important: migration Python files are currently excluded by the repository .gitignore. For team and production deployments, change that policy and commit generated migrations so every environment uses the same schema history.", bold_lead="Important:")

    add_heading(doc, "Smoke Test", 1)
    add_bullets(doc, [
        "Check the API schema at /api/docs/ and confirm the server is reachable.",
        "Register a user, log in, refresh a token, and open the profile endpoint.",
        "Confirm Ollama health and model listing before testing chat.",
        "Create a session, send a prompt, refresh the page, and verify message persistence.",
        "Upload a small PDF, TXT, or DOCX, build the knowledge base, ask a grounded question, and confirm the source list.",
        "Create and reuse a prompt template, then verify dashboard counts and export links.",
    ])

    add_heading(doc, "Production Deployment Plan", 1)
    add_table(doc, ["Component", "Recommended production shape", "Operational requirement"], [
        ("Web API", "Django ASGI served by Uvicorn behind Nginx or a managed reverse proxy", "HTTPS, health checks, timeouts, process restart"),
        ("Frontend", "Build with npm run build and serve static assets from a CDN or Nginx", "Set API base URL and cache immutable assets"),
        ("Database", "PostgreSQL", "Backups, migrations, connection pooling, least-privilege user"),
        ("AI", "Private Ollama host with selected models pre-pulled", "Capacity planning, model health, timeouts, access control"),
        ("Files", "Durable object or volume storage", "Size limits, validation, backup, retention, private downloads"),
        ("Vectors", "Durable Chroma volume or managed vector database", "Backup, user/workspace filters, rebuild procedure"),
        ("Jobs", "Redis plus Celery for indexing and title generation", "Idempotency, retries, status tracking"),
    ], widths=[1.0, 3.3, 2.7])

    add_heading(doc, "Production Checklist", 1)
    add_bullets(doc, [
        "Set DEBUG=False, a strong SECRET_KEY, exact ALLOWED_HOSTS, exact CORS origins, HTTPS, secure proxy headers, and secure cookies.",
        "Commit and apply migrations from a clean database before deployment.",
        "Replace SQLite for multi-user production workloads.",
        "Configure persistent media and vector storage, backups, and restore testing.",
        "Set maximum upload size, allowed extensions, content checks, and private file access.",
        "Add automated tests, lint, build checks, structured logs, and error monitoring.",
        "Set Ollama resource limits and user/workspace rate limits; do not expose Ollama directly to the public internet.",
        "Run a security review of token storage, CORS, CSRF, headers, admin access, permissions, and data deletion behavior.",
    ])

    add_heading(doc, "Troubleshooting", 1)
    add_table(doc, ["Symptom", "Likely cause", "Action"], [
        ("Django import error", "Virtual environment is not active or dependencies are missing", "Activate venv and install requirements"),
        ("Ollama unavailable", "Service is stopped, URL is wrong, or model is missing", "Start Ollama, check OLLAMA_BASE_URL, run ollama list/pull"),
        ("RAG returns no context", "Document is empty, unsupported, or not indexed", "Check file text, build status, and Chroma persistence"),
        ("Scanned PDF has no text", "OCR tools are not installed or paths are unset", "Install Tesseract and Poppler and set environment paths"),
        ("Frontend cannot call API", "Backend or CORS origin mismatch", "Check API URL, server port, and CORS_ALLOWED_ORIGINS"),
        ("Schema errors on a clean install", "Migrations are not committed", "Generate, review, and commit migrations before deployment"),
    ], widths=[1.65, 2.65, 2.7])
    add_heading(doc, "Deployment Ownership Note", 1)
    add_para(doc, "The current repository is a strong local-first development project. Before calling it production-ready, complete the P0 items in the feature inventory: migrations, security settings, tests, persistence, and failure handling. Keep this document versioned beside the code and update it whenever the deployment shape changes.")
    return save(doc, "LocalMind_AI_Deployment_Reference.docx")


if __name__ == "__main__":
    outputs = [feature_inventory(), roadmap(), interview_prep(), deployment_reference()]
    for path in outputs:
        print(path)
